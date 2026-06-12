from flask import Flask, render_template, request, redirect, session
from flask_mysqldb import MySQL
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from io import BytesIO
from flask import send_file
from flask import make_response
import re
from flask import flash
from werkzeug.security import generate_password_hash, check_password_hash
import os
from werkzeug.utils import secure_filename

app = Flask(__name__)
UPLOAD_FOLDER = 'static/profile_pics'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.secret_key = "expense_tracker_secret"

app.config['MYSQL_HOST'] = '127.0.0.1'
app.config['MYSQL_USER'] = 'root'
app.config['MYSQL_PASSWORD'] = 'sumitpatel05'
app.config['MYSQL_DB'] = 'expense_tracker'

mysql = MySQL(app)

@app.route('/')
def home():
    return redirect('/login')

@app.route('/register', methods=['GET', 'POST'])
def register():

    if request.method == 'POST':

        username = request.form['username'].strip()
        email = request.form['email'].strip()
        password = request.form['password']
        security_answer = request.form['security_answer'].strip()

        # Password Validation

        if not re.match(
            r'^(?=.*[a-z])(?=.*[A-Z])(?=.*\d)(?=.*[@$!%*?&]).{8,}$',
            password
        ):

            flash(
                "Password must contain uppercase, lowercase, number and special character",
                "danger"
            )

            return redirect('/register')

        hashed_password = generate_password_hash(password)

        cur = mysql.connection.cursor()

        # Username Check

        cur.execute(
            """
            SELECT *
            FROM users
            WHERE username=%s
            """,
            (username,)
        )

        existing_user = cur.fetchone()

        if existing_user:

            flash(
                "Username already exists",
                "danger"
            )

            cur.close()

            return redirect('/register')

        # Email Check

        cur.execute(
            """
            SELECT *
            FROM users
            WHERE email=%s
            """,
            (email,)
        )

        existing_email = cur.fetchone()

        if existing_email:

            flash(
                "Email already registered",
                "danger"
            )

            cur.close()

            return redirect('/register')

        # Insert User

        cur.execute(
            """
            INSERT INTO users
            (
                username,
                email,
                password,
                security_answer
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s
            )
            """,
            (
                username,
                email,
                hashed_password,
                security_answer
            )
        )

        mysql.connection.commit()

        cur.close()

        flash(
            "Registration Successful",
            "success"
        )

        return redirect('/login')

    return render_template(
        'register.html'
    )

@app.route('/login', methods=['GET', 'POST'])
def login():

    if request.method == 'POST':

        username = request.form['username']
        password = request.form['password']

        cur = mysql.connection.cursor()

        cur.execute(
            """
            SELECT *
            FROM users
            WHERE username=%s
            """,
            (username,)
        )

        user = cur.fetchone()

        cur.close()

        if user and check_password_hash(user[3], password):

            # Save Session
            session['user'] = user[1]
            session['role'] = user[5]

            flash(
                "Login Successful",
                "success"
            )

            # Admin Login
            if user[5] == 'admin':

                return redirect('/admin')

            # Normal User Login
            return redirect('/dashboard')

        else:

            flash(
                "Invalid Username or Password",
                "danger"
            )

            return redirect('/login')

    return render_template(
        'login.html'
    )

@app.route('/dashboard')
def dashboard():

    if 'user' not in session:
        return redirect('/login')

    if session.get('role') == 'admin':
        return redirect('/admin')

    cur = mysql.connection.cursor()

    # User ID

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user = cur.fetchone()

    if not user:
        cur.close()
        return redirect('/login')

    user_id = user[0]

    # Total Expense

    cur.execute(
        "SELECT IFNULL(SUM(amount),0) FROM expenses WHERE user_id=%s",
        (user_id,)
    )

    total_expense = float(cur.fetchone()[0])

    # Top Category

    cur.execute("""
        SELECT category, SUM(amount) as total
        FROM expenses
        WHERE user_id=%s
        GROUP BY category
        ORDER BY total DESC
        LIMIT 1
    """, (user_id,))

    top_category = cur.fetchone()

    # Highest / Lowest / Average

    cur.execute("""
        SELECT
        MAX(amount),
        MIN(amount),
        AVG(amount)
        FROM expenses
        WHERE user_id=%s
    """, (user_id,))

    stats = cur.fetchone()

    highest_expense = stats[0] or 0
    lowest_expense = stats[1] or 0
    average_expense = round(stats[2] or 0, 2)

    # Total Transactions

    cur.execute(
        "SELECT COUNT(*) FROM expenses WHERE user_id=%s",
        (user_id,)
    )

    total_transactions = cur.fetchone()[0]

    # Total Categories

    cur.execute(
        "SELECT COUNT(DISTINCT category) FROM expenses WHERE user_id=%s",
        (user_id,)
    )

    total_categories = cur.fetchone()[0]

    # Pie Chart Data

    cur.execute("""
        SELECT category, SUM(amount)
        FROM expenses
        WHERE user_id=%s
        GROUP BY category
    """, (user_id,))

    chart_data = cur.fetchall()

    # Monthly Expense Data

    cur.execute("""
        SELECT MONTH(expense_date), SUM(amount)
        FROM expenses
        WHERE user_id=%s
        GROUP BY MONTH(expense_date)
        ORDER BY MONTH(expense_date)
    """, (user_id,))

    monthly_data = cur.fetchall()
     
    # Budget

    cur.execute("""
        SELECT monthly_budget
        FROM budgets
        WHERE user_id=%s
        ORDER BY id DESC
        LIMIT 1
    """, (user_id,))

    budget_row = cur.fetchone()

    budget = 0

    if budget_row:
        budget = float(budget_row[0])

    # Remaining Budget

    remaining = budget - total_expense

    # Current Month Expense

    cur.execute("""
     SELECT IFNULL(SUM(amount),0)
     FROM expenses
     WHERE user_id=%s
     AND MONTH(expense_date)=MONTH(CURDATE())
     AND YEAR(expense_date)=YEAR(CURDATE())
     """, (user_id,))

    month_expense = float(cur.fetchone()[0])

    # Last Month Expense

    cur.execute("""
     SELECT IFNULL(SUM(amount),0)
     FROM expenses
     WHERE user_id=%s
     AND MONTH(expense_date)=MONTH(DATE_SUB(CURDATE(), INTERVAL 1 MONTH))
     AND YEAR(expense_date)=YEAR(DATE_SUB(CURDATE(), INTERVAL 1 MONTH))
     """, (user_id,))

    last_month_expense = float(cur.fetchone()[0])

    expense_difference = month_expense - last_month_expense

   # AI Spending Insights

    if total_transactions == 0:

     ai_message = "Add expenses to get AI insights."

    elif budget > 0 and total_expense > budget:

     if top_category:

        ai_message = (
            f"Your highest spending category is {top_category[0]} "
            f"with ₹{top_category[1]:.2f} spent. "
            f"You exceeded your budget by ₹{abs(remaining):.2f}."
        )

     else:

        ai_message = (
            f"You exceeded your budget by ₹{abs(remaining):.2f}."
        )
    
    elif budget > 0 and remaining <= budget * 0.2:

     ai_message = (
        f"Only ₹{remaining:.2f} budget remains. "
        "Spend carefully for the rest of the month."
    )

    elif top_category:

     ai_message = (
        f"Most of your spending is on {top_category[0]}. "
        "Consider reviewing this category."
    )

    else:

     ai_message = "Your spending looks balanced."
    
# Today's Expense

    cur.execute("""
     SELECT IFNULL(SUM(amount),0)
     FROM expenses
     WHERE user_id=%s
     AND expense_date = CURDATE()
     """, (user_id,))

    today_expense = float(cur.fetchone()[0])


   # This Month Expense

    cur.execute("""
     SELECT IFNULL(SUM(amount),0)
     FROM expenses
     WHERE user_id=%s
     AND MONTH(expense_date)=MONTH(CURDATE())
     AND YEAR(expense_date)=YEAR(CURDATE())
     """, (user_id,))

    month_expense = float(cur.fetchone()[0])


# This Week Expense

    cur.execute("""
    SELECT IFNULL(SUM(amount),0)
    FROM expenses
    WHERE user_id=%s
    AND YEARWEEK(expense_date,1)=YEARWEEK(CURDATE(),1)
    """, (user_id,))

    week_expense = float(cur.fetchone()[0])

    # Recent Transactions

    cur.execute("""
        SELECT amount,
        category,
        expense_date
        FROM expenses
        WHERE user_id=%s
        ORDER BY id DESC
        LIMIT 5
        """, (user_id,))

    recent_expenses = cur.fetchall()
    
    cur.execute("""
      SELECT category,
      SUM(amount) as total
      FROM expenses
      WHERE user_id=%s
      GROUP BY category
      ORDER BY total DESC
      LIMIT 5
      """, (user_id,))

    top_categories = cur.fetchall()

    cur.close()

    return render_template(
        'dashboard.html',
        username=session['user'],
        total_expense=total_expense,
        total_transactions=total_transactions,
        total_categories=total_categories,
        chart_data=chart_data,
        monthly_data=monthly_data,
        budget=budget,
        remaining=remaining,
        top_category=top_category,
        highest_expense=highest_expense,
        lowest_expense=lowest_expense,
        average_expense=average_expense,
        recent_expenses=recent_expenses,
        today_expense=today_expense,
        week_expense=week_expense,
        month_expense=month_expense,
        last_month_expense=last_month_expense,
        expense_difference=expense_difference,
        ai_message=ai_message,
        top_categories=top_categories
    )


@app.route('/logout')
def logout():

    session.clear()

    flash(
        "Logged Out Successfully",
        "success"
    )

    return redirect('/login')

@app.route('/add_expense', methods=['GET', 'POST'])
def add_expense():

    if 'user' not in session:
        return redirect('/login')

    if request.method == 'POST':

        try:

            amount = float(request.form['amount'])

            if amount <= 0:

                flash(
                    "Amount must be greater than 0",
                    "danger"
                )

                return redirect('/add_expense')

            category = request.form['category']

            # Custom Category

            if category == "Other":

                custom_category = request.form[
                    'custom_category'
                ].strip()

                if not custom_category:

                    flash(
                        "Please enter custom category",
                        "danger"
                    )

                    return redirect('/add_expense')

                category = custom_category

            description = request.form['description']

            expense_date = request.form['expense_date']

            # Receipt Upload

            receipt_file = request.files['receipt']

            receipt_name = None

            if receipt_file and receipt_file.filename != "":

                receipt_name = secure_filename(
                    receipt_file.filename
                )

                receipt_file.save(
                    os.path.join(
                        "static/receipts",
                        receipt_name
                    )
                )

            cur = mysql.connection.cursor()

            cur.execute(
                """
                SELECT id
                FROM users
                WHERE username=%s
                """,
                (session['user'],)
            )

            user = cur.fetchone()

            user_id = user[0]

            cur.execute(
                """
                INSERT INTO expenses
                (
                    user_id,
                    amount,
                    category,
                    description,
                    expense_date,
                    receipt
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
                """,
                (
                    user_id,
                    amount,
                    category,
                    description,
                    expense_date,
                    receipt_name
                )
            )

            mysql.connection.commit()

            cur.close()

            flash(
                "Expense Added Successfully",
                "success"
            )

            return redirect('/dashboard')

        except Exception as e:

            flash(
                f"Error: {str(e)}",
                "danger"
            )

            return redirect('/add_expense')

    return render_template(
        'add_expense.html'
    )

@app.route('/expenses')
def expenses():

    if 'user' not in session:
        return redirect('/login')

    keyword = request.args.get('keyword', '')
    category = request.args.get('category', '')
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')

    cur = mysql.connection.cursor()

    # Get User ID

    cur.execute(
        """
        SELECT id
        FROM users
        WHERE username=%s
        """,
        (session['user'],)
    )

    user = cur.fetchone()

    user_id = user[0]

    # Base Query

    query = """
        SELECT
            id,
            amount,
            category,
            description,
            expense_date,
            receipt
        FROM expenses
        WHERE user_id=%s
    """

    values = [user_id]

    # Search Description

    if keyword:

        query += """
            AND description LIKE %s
        """

        values.append(
            f"%{keyword}%"
        )

    # Category Filter

    if category:

        query += """
            AND category=%s
        """

        values.append(
            category
        )

    # Date Filter

    if start_date and end_date:

        query += """
            AND expense_date
            BETWEEN %s AND %s
        """

        values.append(
            start_date
        )

        values.append(
            end_date
        )

    # Latest First

    query += """
        ORDER BY expense_date DESC
    """

    cur.execute(
        query,
        tuple(values)
    )

    expenses = cur.fetchall()

    cur.close()

    return render_template(
        'expenses.html',
        expenses=expenses
    )

@app.route('/edit_expense/<int:id>', methods=['GET','POST'])
def edit_expense(id):

    if 'user' not in session:
        return redirect('/login')

    cur = mysql.connection.cursor()

    if request.method == 'POST':

        amount = request.form['amount']
        category = request.form['category']
        description = request.form['description']
        expense_date = request.form['expense_date']

        cur.execute(
            """
            UPDATE expenses
            SET amount=%s,
                category=%s,
                description=%s,
                expense_date=%s
            WHERE id=%s
            """,
            (amount, category, description, expense_date, id)
        )

        mysql.connection.commit()

        cur.close()

        flash("Expense Updated Successfully", "success")

        return redirect('/expenses')

    cur.execute(
        "SELECT * FROM expenses WHERE id=%s",
        (id,)
    )

    expense = cur.fetchone()

    cur.close()

    return render_template(
        'edit_expense.html',
        expense=expense
    )

@app.route('/delete_expense/<int:id>')
def delete_expense(id):

    if 'user' not in session:
        return redirect('/login')

    cur = mysql.connection.cursor()

    cur.execute(
        "DELETE FROM expenses WHERE id=%s",
        (id,)
    )

    mysql.connection.commit()

    cur.close()

    flash("Expense Deleted Successfully", "success")

    return redirect('/expenses')

@app.route('/budget', methods=['GET', 'POST'])
def budget():

    if 'user' not in session:
        return redirect('/login')

    cur = mysql.connection.cursor()

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user_id = cur.fetchone()[0]

    if request.method == 'POST':

        budget_amount = request.form['budget']

        # Check existing budget

        cur.execute(
            "SELECT * FROM budgets WHERE user_id=%s",
            (user_id,)
        )

        existing_budget = cur.fetchone()

        if existing_budget:

            # Update Budget

            cur.execute(
                """
                UPDATE budgets
                SET monthly_budget=%s
                WHERE user_id=%s
                """,
                (budget_amount, user_id)
            )

        else:

            # Insert Budget

            cur.execute(
                """
                INSERT INTO budgets(user_id, monthly_budget)
                VALUES(%s,%s)
                """,
                (user_id, budget_amount)
            )

        mysql.connection.commit()

        cur.close()

        return redirect('/dashboard')

    cur.close()

    return render_template('budget.html')

@app.route('/download_report')
def download_report():

    if 'user' not in session:
        return redirect('/login')

    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer
    )
    from reportlab.lib.styles import getSampleStyleSheet
    from io import BytesIO
    from flask import send_file
    from datetime import datetime

    cur = mysql.connection.cursor()

    # User ID

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user_id = cur.fetchone()[0]

    # Expenses

    cur.execute("""
        SELECT
        category,
        amount,
        description,
        expense_date
        FROM expenses
        WHERE user_id=%s
        ORDER BY expense_date DESC
    """, (user_id,))

    expenses = cur.fetchall()

    # Total Expense

    cur.execute("""
        SELECT IFNULL(SUM(amount),0)
        FROM expenses
        WHERE user_id=%s
    """, (user_id,))

    total_expense = float(cur.fetchone()[0])

    cur.close()

    # PDF Buffer

    buffer = BytesIO()

    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()

    elements = []

    # Title

    title = Paragraph(
        "AI Smart Expense Tracker Report",
        styles['Title']
    )

    elements.append(title)

    elements.append(Spacer(1, 12))

    # User Info

    elements.append(
        Paragraph(
            f"<b>Username:</b> {session['user']}",
            styles['Normal']
        )
    )

    elements.append(
        Paragraph(
            f"<b>Generated On:</b> {datetime.now().strftime('%d-%m-%Y %H:%M')}",
            styles['Normal']
        )
    )

    elements.append(
        Paragraph(
            f"<b>Total Expense:</b> ₹ {total_expense}",
            styles['Normal']
        )
    )

    elements.append(
        Paragraph(
            f"<b>Total Transactions:</b> {len(expenses)}",
            styles['Normal']
        )
    )

    elements.append(Spacer(1, 20))

    # Table Data

    data = [
        [
            "Category",
            "Amount",
            "Description",
            "Date"
        ]
    ]

    for expense in expenses:

        data.append([
            expense[0],
            f"₹ {expense[1]}",
            expense[2] or "",
            str(expense[3])
        ])

    table = Table(
        data,
        colWidths=[100, 80, 180, 100]
    )

    table.setStyle(
        TableStyle([

            ('BACKGROUND', (0,0), (-1,0), colors.darkblue),

            ('TEXTCOLOR', (0,0), (-1,0), colors.white),

            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),

            ('GRID', (0,0), (-1,-1), 1, colors.black),

            ('BACKGROUND', (0,1), (-1,-1), colors.whitesmoke),

            ('ALIGN', (1,1), (1,-1), 'CENTER'),

        ])
    )

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="AI_Expense_Report.pdf",
        mimetype="application/pdf"
    )

@app.route('/search')
def search():

    if 'user' not in session:
        return redirect('/login')

    keyword = request.args.get('keyword', '')

    cur = mysql.connection.cursor()

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user_id = cur.fetchone()[0]

    cur.execute("""
        SELECT *
        FROM expenses
        WHERE user_id=%s
        AND category LIKE %s
    """, (user_id, f"%{keyword}%"))

    expenses = cur.fetchall()

    return render_template(
        'expenses.html',
        expenses=expenses
    )
@app.route('/filter')
def filter_expenses():

    if 'user' not in session:
        return redirect('/login')

    start = request.args.get('start_date')
    end = request.args.get('end_date')

    cur = mysql.connection.cursor()

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user_id = cur.fetchone()[0]

    # Agar date select nahi ki gayi
    if not start or not end:

        cur.execute(
            "SELECT * FROM expenses WHERE user_id=%s",
            (user_id,)
        )

    else:

        cur.execute("""
            SELECT *
            FROM expenses
            WHERE user_id=%s
            AND expense_date BETWEEN %s AND %s
        """, (user_id, start, end))

    expenses = cur.fetchall()

    cur.close()

    return render_template(
        'expenses.html',
        expenses=expenses
    )
@app.route('/export_excel')
def export_excel():

    if 'user' not in session:
        return redirect('/login')

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from io import BytesIO
    from flask import send_file
    from datetime import datetime

    cur = mysql.connection.cursor()

    # User ID

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user_id = cur.fetchone()[0]

    # Expenses

    cur.execute("""
        SELECT amount,
               category,
               description,
               expense_date
        FROM expenses
        WHERE user_id=%s
        ORDER BY expense_date DESC
    """, (user_id,))

    expenses = cur.fetchall()

    # Total Expense

    cur.execute("""
        SELECT IFNULL(SUM(amount),0)
        FROM expenses
        WHERE user_id=%s
    """, (user_id,))

    total_expense = float(cur.fetchone()[0])

    wb = Workbook()
    ws = wb.active

    ws.title = "Expense Report"

    # Report Title

    ws.merge_cells('A1:D1')

    title = ws['A1']

    title.value = "AI Smart Expense Tracker Report"

    title.font = Font(
        bold=True,
        size=18
    )

    title.alignment = Alignment(
        horizontal='center'
    )

    # Username

    ws['A3'] = "Username"
    ws['B3'] = session['user']

    # Date

    ws['A4'] = "Generated On"
    ws['B4'] = datetime.now().strftime("%d-%m-%Y %H:%M")

    # Total Expense

    ws['A5'] = "Total Expense"
    ws['B5'] = f"₹ {total_expense}"

    # Total Transactions

    ws['A6'] = "Total Transactions"
    ws['B6'] = len(expenses)

    # Headers

    headers = [
        "Amount",
        "Category",
        "Description",
        "Date"
    ]

    header_fill = PatternFill(
        start_color="4F81BD",
        end_color="4F81BD",
        fill_type="solid"
    )

    header_row = 8

    for col_num, header in enumerate(headers, 1):

        cell = ws.cell(
            row=header_row,
            column=col_num
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

    # Data

    for row_num, expense in enumerate(
        expenses,
        start=9
    ):

        ws.cell(
            row=row_num,
            column=1,
            value=expense[0]
        )

        ws.cell(
            row=row_num,
            column=2,
            value=expense[1]
        )

        ws.cell(
            row=row_num,
            column=3,
            value=expense[2]
        )

        ws.cell(
            row=row_num,
            column=4,
            value=expense[3]
        )

    # Column Width

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 20

    file_stream = BytesIO()

    wb.save(file_stream)

    file_stream.seek(0)

    cur.close()

    return send_file(
        file_stream,
        as_attachment=True,
        download_name="AI_Expense_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/profile')
def profile():

    if 'user' not in session:
        return redirect('/login')

    cur = mysql.connection.cursor()

    # User Details
    cur.execute(
    """
    SELECT id, username, email, profile_pic
    FROM users
    WHERE username=%s
    """,
    (session['user'],)
     )

    user = cur.fetchone()

    user_id = user[0]

    # Total Expense
    cur.execute(
        """
        SELECT IFNULL(SUM(amount),0)
        FROM expenses
        WHERE user_id=%s
        """,
        (user_id,)
    )

    total_expense = float(cur.fetchone()[0])

    # Total Transactions
    cur.execute(
        """
        SELECT COUNT(*)
        FROM expenses
        WHERE user_id=%s
        """,
        (user_id,)
    )

    total_transactions = cur.fetchone()[0]

    # Budget
    cur.execute(
        """
        SELECT monthly_budget
        FROM budgets
        WHERE user_id=%s
        ORDER BY id DESC
        LIMIT 1
        """,
        (user_id,)
    )

    budget_row = cur.fetchone()

    budget = 0

    if budget_row:
        budget = float(budget_row[0])

    cur.close()

    return render_template(
        'profile.html',
        user=user,
        total_expense=total_expense,
        total_transactions=total_transactions,
        budget=budget
    )

@app.route('/change_password', methods=['POST'])
def change_password():

    if 'user' not in session:
        return redirect('/login')

    current_password = request.form['current_password']
    new_password = request.form['new_password']
    confirm_password = request.form['confirm_password']

    # Confirm Password Check
    if new_password != confirm_password:
        return """
        <script>
        alert('New Password and Confirm Password do not match');
        window.location='/profile';
        </script>
        """

    # Current Password Check
    cur = mysql.connection.cursor()

    cur.execute(
        """
        SELECT password
        FROM users
        WHERE username=%s
        """,
        (session['user'],)
    )

    db_password = cur.fetchone()[0]

    if current_password != db_password:
        return """
        <script>
        alert('Current Password Incorrect');
        window.location='/profile';
        </script>
        """

    # Strong Password Validation
    password_pattern = (
        r'^(?=.*[a-z])'
        r'(?=.*[A-Z])'
        r'(?=.*\d)'
        r'(?=.*[@$!%*?&])'
        r'[A-Za-z\d@$!%*?&]{8,}$'
    )

    if not re.match(password_pattern, new_password):
        return """
        <script>
        alert('Password must contain:\\n\\nMinimum 8 characters\\n1 Capital Letter\\n1 Small Letter\\n1 Number\\n1 Special Character');
        window.location='/profile';
        </script>
        """

    # Update Password
    cur.execute(
        """
        UPDATE users
        SET password=%s
        WHERE username=%s
        """,
        (new_password, session['user'])
    )

    mysql.connection.commit()

    cur.close()

    return """
    <script>
    alert('Password Changed Successfully');
    window.location='/profile';
    </script>
    """

@app.route('/delete_account')
def delete_account():

    if 'user' not in session:
        return redirect('/login')

    cur = mysql.connection.cursor()

    cur.execute(
        "SELECT id FROM users WHERE username=%s",
        (session['user'],)
    )

    user = cur.fetchone()

    if user:

        user_id = user[0]

        # Delete Expenses

        cur.execute(
            "DELETE FROM expenses WHERE user_id=%s",
            (user_id,)
        )

        # Delete Budget

        cur.execute(
            "DELETE FROM budgets WHERE user_id=%s",
            (user_id,)
        )

        # Delete User

        cur.execute(
            "DELETE FROM users WHERE id=%s",
            (user_id,)
        )

        mysql.connection.commit()

    cur.close()

    session.clear()

    flash(
        "Account Deleted Successfully",
        "success"
    )

    return redirect('/register')
@app.route('/toggle_theme')
def toggle_theme():

    if 'theme' not in session:

        session['theme'] = 'dark'

    if session['theme'] == 'dark':

        session['theme'] = 'light'

    else:

        session['theme'] = 'dark'

    return redirect(request.referrer)

@app.route('/admin')
def admin():

    if 'user' not in session:
        return redirect('/login')

    if session.get('role') != 'admin':

        flash(
            "Access Denied",
            "danger"
        )

        return redirect('/dashboard')

    cur = mysql.connection.cursor()

    # Total Users

    cur.execute(
        "SELECT COUNT(*) FROM users"
    )

    total_users = cur.fetchone()[0]

    # Total Expenses

    cur.execute(
        "SELECT COUNT(*) FROM expenses"
    )

    total_expenses = cur.fetchone()[0]

    # Total Amount

    cur.execute(
        """
        SELECT IFNULL(SUM(amount),0)
        FROM expenses
        """
    )

    total_amount = float(cur.fetchone()[0])

    # Total Admins

    cur.execute(
        """
        SELECT COUNT(*)
        FROM users
        WHERE role='admin'
        """
    )

    total_admins = cur.fetchone()[0]

    # Total Users (Normal)

    cur.execute(
        """
        SELECT COUNT(*)
        FROM users
        WHERE role='user'
        """
    )

    total_normal_users = cur.fetchone()[0]

    # Average Expense

    cur.execute(
        """
        SELECT IFNULL(AVG(amount),0)
        FROM expenses
        """
    )

    average_expense = round(
        float(cur.fetchone()[0]),
        2
    )

    # Users List

    cur.execute(
        """
        SELECT
            id,
            username,
            email,
            role
        FROM users
        ORDER BY id DESC
        """
    )

    users = cur.fetchall()

    cur.close()

    return render_template(
        'admin.html',
        total_users=total_users,
        total_expenses=total_expenses,
        total_amount=total_amount,
        total_admins=total_admins,
        total_normal_users=total_normal_users,
        average_expense=average_expense,
        users=users
    )

@app.route('/delete_user/<int:user_id>')
def delete_user(user_id):

    if 'user' not in session:
        return redirect('/login')

    if session.get('role') != 'admin':

        flash(
            "Access Denied",
            "danger"
        )

        return redirect('/dashboard')

    cur = mysql.connection.cursor()

    # Admin khud ko delete na kar sake

    cur.execute(
        "SELECT username FROM users WHERE id=%s",
        (user_id,)
    )

    user = cur.fetchone()

    if user and user[0] == session['user']:

        flash(
            "You cannot delete yourself",
            "danger"
        )

        cur.close()

        return redirect('/admin')

    # User Expenses Delete

    cur.execute(
        "DELETE FROM expenses WHERE user_id=%s",
        (user_id,)
    )

    # Budget Delete

    cur.execute(
        "DELETE FROM budgets WHERE user_id=%s",
        (user_id,)
    )

    # User Delete

    cur.execute(
        "DELETE FROM users WHERE id=%s",
        (user_id,)
    )

    mysql.connection.commit()

    cur.close()

    flash(
        "User Deleted Successfully",
        "success"
    )

    return redirect('/admin')

@app.route('/forgot_password', methods=['GET', 'POST'])
def forgot_password():

    if request.method == 'POST':

        username = request.form['username'].strip()
        answer = request.form['security_answer'].strip()
        new_password = request.form['new_password']

        cur = mysql.connection.cursor()

        cur.execute(
            """
            SELECT security_answer
            FROM users
            WHERE username=%s
            """,
            (username,)
        )

        user = cur.fetchone()

        db_answer = ""

        if user and user[0]:

            db_answer = str(user[0]).strip().lower()

        entered_answer = answer.strip().lower()

        print("Database Answer =", db_answer)
        print("Entered Answer =", entered_answer)

        if db_answer == entered_answer:

            hashed_password = generate_password_hash(
                new_password
            )

            cur.execute(
                """
                UPDATE users
                SET password=%s
                WHERE username=%s
                """,
                (
                    hashed_password,
                    username
                )
            )

            mysql.connection.commit()

            flash(
                "Password Reset Successful",
                "success"
            )

            cur.close()

            return redirect('/login')

        else:

            flash(
                "Invalid Security Answer",
                "danger"
            )

            cur.close()

            return redirect('/forgot_password')

    return render_template(
        'forgot_password.html'
    )
if __name__ == '__main__':
    app.run(debug=True)